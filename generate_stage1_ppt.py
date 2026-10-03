# -*- coding: utf-8 -*-
"""Generate the stage-1 seminar PPT draft for the Mandol constrained-retrieval task.

Run:
    uv run --no-sync python generate_stage1_ppt.py

Output:
    阶段1研讨-PPT初稿.pptx  (16:9, 15 content slides + 5 backup slides)

Notes:
    - Edit NAME / DATE below, then re-run to refresh the deck.
    - Slide content mirrors 阶段1研讨PPT大纲.md and 系统分析与设计报告.md.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------------------------------------------------------------- config

FONT = "Microsoft YaHei"

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x75, 0xB6)
SKY = RGBColor(0x9D, 0xC3, 0xE6)
LIGHT = RGBColor(0xDE, 0xEB, 0xF7)
ICE = RGBColor(0xF2, 0xF7, 0xFC)
ACCENT = RGBColor(0xC5, 0x5A, 0x11)
GREEN = RGBColor(0x3E, 0x7A, 0x35)
GRAY = RGBColor(0x3F, 0x3F, 0x3F)
MIDGRAY = RGBColor(0x80, 0x80, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

NAME = "（填写姓名）"
DATE = "2026-10"
OUTPUT = "阶段1研讨-PPT初稿.pptx"

TOTAL = 20  # 15 content + 5 backup

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

X0 = 0.6          # left margin
FULL = 12.13      # full content width
TOTAL_W = 13.333


# ---------------------------------------------------------------- helpers

def _apply_font(run, size, bold, color, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", FONT)


def add_box(slide, x, y, w, h):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    return box, tf


def para(tf, first, text, size=14, bold=False, color=GRAY,
         align=PP_ALIGN.LEFT, space_after=6, space_before=0,
         level=0, line_spacing=1.15):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    p.line_spacing = line_spacing
    p.level = level
    run = p.add_run()
    run.text = text
    _apply_font(run, size, bold, color)
    return p


def para_runs(tf, first, parts, size=14, align=PP_ALIGN.LEFT, space_after=6,
              space_before=0, level=0, line_spacing=1.15):
    """parts: [(text, bold, color), ...]"""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    p.line_spacing = line_spacing
    p.level = level
    for text, bold, color in parts:
        run = p.add_run()
        run.text = text
        _apply_font(run, size, bold, color)
    return p


def add_shape(slide, shape_type, x, y, w, h, fill, line=None, radius=None):
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    if radius is not None and shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    return shp


def chip(slide, x, y, w, h, text, sub=None, fill=LIGHT, color=NAVY,
         size=12.5, sub_size=10, radius=0.22, line=None):
    shp = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h,
                    fill, line=line, radius=radius)
    tf = shp.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, True, text, size=size, bold=True, color=color,
         align=PP_ALIGN.CENTER, space_after=0 if sub else 0)
    if sub:
        para(tf, False, sub, size=sub_size, bold=False, color=color,
             align=PP_ALIGN.CENTER, space_after=0, space_before=1)
    return shp


def arrow(slide, x, y, w, h, color=BLUE):
    return add_shape(slide, MSO_SHAPE.RIGHT_ARROW, x, y, w, h, color)


def band(slide, x, y, w, h, text, fill=NAVY, color=WHITE, size=13.5, bold=True):
    """Full-width highlight band with centered text."""
    shp = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill,
                    radius=0.14)
    tf = shp.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, True, text, size=size, bold=bold, color=color,
         align=PP_ALIGN.CENTER, space_after=0)
    return shp


def header(slide, title, page, subtitle=None, backup=False):
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(X0 - 0.05), Inches(0.44),
              Inches(0.1), Inches(0.5), ACCENT if backup else BLUE)
    box, tf = add_box(slide, Inches(X0 + 0.18), Inches(0.34), Inches(FULL - 2.0), Inches(0.72))
    para(tf, True, title, size=22, bold=True, color=NAVY, space_after=0)
    if subtitle:
        para(tf, False, subtitle, size=10.5, color=MIDGRAY, space_after=0, space_before=1)
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(X0), Inches(1.2),
              Inches(FULL), Pt(1.2), LIGHT)
    # footer
    box, tf = add_box(slide, Inches(X0), Inches(7.08), Inches(9.0), Inches(0.3))
    para(tf, True, "Mandol 融合索引增强 ｜ 阶段 1 课堂研讨" + (" ｜ 备用页" if backup else ""),
         size=9, bold=False, color=MIDGRAY, space_after=0)
    box, tf = add_box(slide, Inches(TOTAL_W - 1.4), Inches(7.08), Inches(0.85), Inches(0.3))
    para(tf, True, f"{page} / {TOTAL}", size=9, bold=False, color=MIDGRAY,
         align=PP_ALIGN.RIGHT, space_after=0)


def add_table(slide, x, y, w, data, col_widths, font_size=11, header_size=None,
              row_height=0.42, first_col_bold=False, highlight_rows=(),
              highlight_fill=LIGHT):
    rows, cols = len(data), len(data[0])
    shape = slide.shapes.add_table(rows, cols, Inches(x), Inches(y),
                                   Inches(w), Inches(row_height * rows))
    table = shape.table
    table.first_row = False
    table.horz_banding = False
    for c, cw in enumerate(col_widths):
        table.columns[c].width = Inches(cw)
    for r in range(rows):
        table.rows[r].height = Inches(row_height)
        for c in range(cols):
            cell = table.cell(r, c)
            cell.margin_left = Inches(0.07)
            cell.margin_right = Inches(0.07)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.space_after = Pt(0)
            p.line_spacing = 1.05
            run = p.add_run()
            run.text = str(data[r][c])
            if r == 0:
                _apply_font(run, header_size or font_size, True, WHITE)
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY
            else:
                is_hl = r in highlight_rows
                bold = (first_col_bold and c == 0) or is_hl
                _apply_font(run, font_size, bold, NAVY if bold else GRAY)
                cell.fill.solid()
                if is_hl:
                    cell.fill.fore_color.rgb = highlight_fill
                else:
                    cell.fill.fore_color.rgb = WHITE if r % 2 == 1 else ICE
    return table


def chips_row(slide, y, items, h=0.62, fill=LIGHT, color=NAVY, size=12.5,
              margin=X0, full=FULL, gap=0.2, radius=0.22):
    n = len(items)
    w = (full - gap * (n - 1)) / n
    for i, item in enumerate(items):
        x = margin + i * (w + gap)
        if isinstance(item, tuple):
            chip(slide, x, y, w, h, item[0], sub=item[1], fill=fill,
                 color=color, size=size, radius=radius)
        else:
            chip(slide, x, y, w, h, item, fill=fill, color=color,
                 size=size, radius=radius)


def flow_row(slide, y, h, items, margin=X0, full=FULL, arrow_w=0.3,
             fill=LIGHT, color=NAVY, size=12.5, sub_size=10):
    """items: list of (text, sub). Draws chip -> chip -> ... horizontally."""
    n = len(items)
    chip_w = (full - arrow_w * (n - 1)) / n
    x = margin
    for i, (text, sub) in enumerate(items):
        chip(slide, x, y, chip_w, h, text, sub=sub, fill=fill, color=color,
             size=size, sub_size=sub_size)
        x += chip_w
        if i < n - 1:
            arrow(slide, x + 0.04, y + h / 2 - 0.11, arrow_w - 0.08, 0.22)
            x += arrow_w


def new_slide():
    return prs.slides.add_slide(BLANK)


# ================================================================ P1 cover

s = new_slide()
add_shape(s, MSO_SHAPE.RECTANGLE, 0, 0, TOTAL_W, 7.5, NAVY)
add_shape(s, MSO_SHAPE.RECTANGLE, 0, 4.72, TOTAL_W, 0.03, BLUE)
box, tf = add_box(s, 0.95, 1.7, 11.6, 2.0)
para(tf, True, "Mandol 融合索引增强", size=40, bold=True, color=WHITE, space_after=8)
para(tf, False, "组合查询的问题分析与初步方案", size=24, color=SKY, space_after=0)
box, tf = add_box(s, 0.95, 3.85, 11.6, 0.5)
para(tf, True, "约束下推式组合检索（Constrained Hybrid Retrieval）",
     size=16, bold=False, color=RGBColor(0x8E, 0xA9, 0xDB), space_after=0)
box, tf = add_box(s, 0.95, 5.15, 11.6, 1.3)
para(tf, True, "《大模型辅助的面向对象软件分析、设计与开发》课程大作业 ｜ 阶段 1 课堂研讨",
     size=13, bold=False, color=LIGHT, space_after=6)
para(tf, False, f"汇报人：{NAME}      {DATE}", size=12,
     color=RGBColor(0x8E, 0xA9, 0xDB), space_after=0)


# ================================================================ P2 background

s = new_slide()
header(s, "任务背景：我们要解决什么", 2)
box, tf = add_box(s, X0, 1.42, FULL, 3.3)
para_runs(tf, True, [("课程方向：", True, NAVY),
                     ("融合索引增强 —— 让记忆访问支持“语义相似 + 精确标识 + 属性约束 + 关系查询”的协同", False, GRAY)],
          size=15, space_after=10)
para(tf, False, "目标场景（作业样例）：", size=14, bold=True, color=NAVY, space_after=4)
para(tf, False, "“与给定内容语义相近、发生于某时间之后，并且属于某实体邻居的记忆。”",
     size=16, bold=True, color=ACCENT, space_after=12, level=1)
para_runs(tf, False, [("Mandol ：", True, NAVY),
                      ("面向大模型智能体的记忆组织、存储与检索系统", False, GRAY)],
          size=14, space_after=6)
para(tf, False, "本方向要求扩展多种访问条件的协同能力，并比较检索效果、访问开销与维护成本。",
     size=13, color=MIDGRAY, space_after=0)
chips_row(s, 5.55, ["分层记忆模型", "融合存储（KV + 向量 + 图）", "主动检索管线", "分层持久化"],
          h=0.7, size=13.5)


# ================================================================ P3 system tour

s = new_slide()
header(s, "系统速览：检索管线与关键对象", 3)
flow_row(s, 1.5, 1.0, [
    ("QueryBundle", "查询编码缓存"),
    ("三路召回", "BM25 / SPLADE / Cosine"),
    ("分数融合", "RRF / 加权 / MMR"),
    ("重排", "Reranker"),
    ("图扩展", "GraphContextExpander"),
])
box, tf = add_box(s, X0, 2.85, FULL, 0.9)
para_runs(tf, True, [("关键对象：", True, NAVY),
                     ("MemoryUnit（原子容器）｜ SemanticMap（融合存储主类）｜ SemanticGraph（结构关系层）｜ BaseRetriever 族（多态扩展点）",
                      False, GRAY)], size=13.5, space_after=0)
box, tf = add_box(s, X0, 3.75, FULL, 1.2)
para_runs(tf, True, [("写入链的天然挂载点：", True, NAVY),
                     ("FAISS 增量、BM25/SPLADE 增量（_incremental_aux_retriever_add）、换页触发 —— 新增索引与它们并列挂载",
                      False, GRAY)], size=13.5, space_after=6)
para_runs(tf, False, [("检索链的现成缝：", True, ACCENT),
                      ("三路检索器均已原生支持 candidate_uids 过滤与 space_names 参数（抽象基类 **kwargs 透传）",
                       False, GRAY)], size=13.5, space_after=0)
band(s, X0, 5.45, FULL, 0.85,
     "关键发现：候选集下推链路已贯通 —— 方案可零修改接入，不必新增召回源")


# ================================================================ P4 problem

s = new_slide()
header(s, "问题聚焦：组合查询“无路可走”", 4)
box = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, X0, 1.4, FULL, 0.8, ICE,
                line=BLUE, radius=0.14)
tf = box.text_frame
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
para(tf, True, "目标查询（U3）：与内容 C 语义相近 ＋ 发生在 T 之后 ＋ 属于实体 E 的 1-hop 邻居 ＋ 属于空间 S",
     size=13.5, bold=True, color=NAVY, align=PP_ALIGN.CENTER, space_after=0)

blocks = [
    ("① 先语义检索再手工过滤", "top_k 失真；结果不保证硬满足全部约束", "过滤只能在后，无法参与召回"),
    ("② filter_memory_units 全库线性扫描", "高选择性场景代价大；无索引、无分页", "游离于 BaseRetriever 多态体系之外"),
    ("③ 图查询与语义检索平行", "“实体邻居”无法作为附加约束参与组合", "图证据进不了组合查询"),
]
bw, bgap = 3.85, 0.29
for i, (t, m1, m2) in enumerate(blocks):
    x = X0 + i * (bw + bgap)
    box = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.5, bw, 2.75, ICE, radius=0.08)
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_top = Inches(0.14)
    para(tf, True, t, size=13.5, bold=True, color=NAVY, space_after=8)
    para(tf, False, m1, size=12, color=GRAY, space_after=6)
    para(tf, False, m2, size=11, color=MIDGRAY, space_after=0)
band(s, X0, 5.75, FULL, 0.8,
     "现有范式 = “并集 + 分数融合”   ｜   组合查询需要 = “交集 + 谓词下推”")


# ================================================================ P5 requirements

s = new_slide()
header(s, "需求分析：四条可验证需求 + 三个用例", 5)
add_table(s, X0, 1.5, FULL, [
    ["需求", "含义", "验收方式"],
    ["R1 表达能力", "一条语句表达“语义 + 属性 + 时间 + 关系”组合约束", "单次调用完成（而非分步手工过滤）"],
    ["R2 正确性", "结果硬满足全部约束（不是“尽量”）", "集成测试断言 + 与线性过滤基线对拍"],
    ["R3 效率", "高选择性约束避免全库扫描与无效向量比对", "pre-filter 下推；候选占比扫描曲线"],
    ["R4 衔接与成本", "不破坏现有流程；维护成本与收益可量化", "回归测试零修改通过；构建/增量/内存指标"],
], col_widths=[2.3, 6.0, 3.83], font_size=11.5, header_size=12, row_height=0.62)
chips_row(s, 5.15, ["U1 语义 + 时间", "U2 语义 + 关系", "U3 全组合（主用例）"], h=0.62, size=13)
box, tf = add_box(s, X0, 6.05, FULL, 0.6)
para(tf, True, "U3 覆盖全部四类条件：“与 C 语义相近 + 发生在 T 之后 + 属于实体 E 的 1-hop 邻居 + 属于空间 S”",
     size=11.5, bold=False, color=MIDGRAY, space_after=0)


# ================================================================ P6 capability map

s = new_slide()
header(s, "缺口分析（一）：能力地图", 6)
add_table(s, X0, 1.5, FULL, [
    ["访问形态", "现成能力", "缺口"],
    ["语义相似", "三路召回 + RRF/加权/MMR 融合 + 重排", "—"],
    ["精确标识", "uid 精确存取、空间成员过滤（candidate_uids 子集检索）", "—"],
    ["属性约束", "filter_memory_units 线性过滤（10 操作符）", "游离多态体系之外；不查 metadata；无索引、无分页"],
    ["时间约束", "时间戳随 payload 存储（raw_data[“timestamp”]）", "字符串、格式异构、无规范化、无有序索引"],
    ["关系查询", "GraphRetriever 节点/边检索；图上 BFS 能力", "与语义检索平行，不能作为附加约束；部分图 API 断链"],
    ["组合查询", "无", "没有“交集 + 谓词下推”的表达与调度机制"],
], col_widths=[1.9, 5.4, 4.83], font_size=11, header_size=12, row_height=0.52,
    highlight_rows={6})
band(s, X0, 5.35, FULL, 0.85,
     "底层 pre-filter 管道已打通 ｜ 缺失的是：候选集生成层（属性/时间/关系）＋ 调度层（协调与路径选择）",
     size=13)


# ================================================================ P7 three gaps

s = new_slide()
header(s, "缺口分析（二）：三个具体缺口（证据级）", 7)
add_table(s, X0, 1.5, FULL, [
    ["缺口", "源码事实（核验）", "证据"],
    ["属性约束", "filter_memory_units：10 操作符但为线性扫描；取值 getattr/raw_data、不查 metadata；无分页",
     "semantic_map.py:3169"],
    ["时间约束", "时间戳在 raw_data[“timestamp”]，字符串且格式随数据集而异；created_time 写入不填充、仅序列化恢复",
     "locomo_benchmark_episodic.py:462\nsemantic_map.py:3054"],
    ["关系约束", "“实体邻居”无下推路径；既有图 API 内部依赖未定义方法（全仓库仅有调用点），运行时必抛异常",
     "semantic_graph.py:1161 / 1126-1132 / 1204"],
], col_widths=[1.9, 7.6, 2.63], font_size=11, header_size=12, row_height=0.95,
    first_col_bold=True)
box, tf = add_box(s, X0, 5.7, FULL, 0.9)
para(tf, True, "核验方法：大模型建议 → 定位文件/行 → 静态核验（调用链/签名/数据结构）→ 采纳 / 修正 / 弃用",
     size=12.5, bold=False, color=GRAY, space_after=4)
para(tf, False, "※ 核验故事与完整对照表见 P13 与备用页 B1；行号级证据索引见备用页 B4。",
     size=11, color=MIDGRAY, space_after=0)


# ================================================================ P8 solution overview

s = new_slide()
header(s, "初步方案总览：约束下推式组合检索", 8)
flow_row(s, 1.5, 1.05, [
    ("QueryConstraints", "声明式约束（值对象）"),
    ("Resolver 策略族", "各叶子谓词独立求解"),
    ("CandidateSet", "候选集代数（AND/OR/NOT）"),
    ("Planner", "选择率自适应调度"),
    ("smart_search", "candidate_uids 下推"),
], sub_size=10.5)
box, tf = add_box(s, X0, 2.95, FULL, 2.3)
para_runs(tf, True, [("核心叙事：", True, ACCENT),
                     ("本方案的实质不是“多加几个索引”，而是基于选择率估计的“自适应执行计划”——借鉴数据库查询优化的经典范式。",
                      False, GRAY)], size=14.5, space_after=10)
para_runs(tf, False, [("· ", False, BLUE), ("倒排索引 / 时间有序索引 / 图 BFS 都只是“候选集生成器”角色下的可替换实现。", False, GRAY)],
          size=13.5, space_after=6)
para_runs(tf, False, [("· ", False, BLUE), ("复用现有 BaseRetriever 多态体系，以协作类接入，底层三路检索器零修改。", False, GRAY)],
          size=13.5, space_after=6)
para_runs(tf, False, [("· ", False, BLUE), ("语义相似度只作为排序信号最后介入，保证重排不浪费在被过滤掉的候选上。", False, GRAY)],
          size=13.5, space_after=0)
band(s, X0, 5.75, FULL, 0.8,
     "先穷尽既有缝，再考虑开新口子 —— candidate_uids 就是那条缝", size=13.5)


# ================================================================ P9 object model

s = new_slide()
header(s, "对象建模：双轨多态", 9)
bw = 3.75
gap = 0.44
y, h = 1.5, 3.6
x1 = X0
x2 = X0 + bw + gap
x3 = X0 + 2 * (bw + gap)

box = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x1, y, bw, h, ICE, radius=0.06)
tf = box.text_frame
tf.margin_top = Inches(0.12)
para(tf, True, "约束解析多态", size=14.5, bold=True, color=NAVY, space_after=2)
para(tf, False, "BaseConstraintResolver（ABC）策略族", size=10.5, color=MIDGRAY, space_after=8)
for t in ["MetadataConstraintResolver（属性）", "TimeRangeConstraintResolver（时间）",
          "RelationConstraintResolver（关系 BFS）", "SpaceConstraintResolver（空间）"]:
    para(tf, False, "· " + t, size=11.5, color=GRAY, space_after=7)

box = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x2, y + 0.55, bw, h - 1.1, LIGHT, radius=0.08)
tf = box.text_frame
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
para(tf, True, "CandidateSet", size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER, space_after=2)
para(tf, False, "值对象：uid 集合 + 选择性估计", size=11, color=GRAY, align=PP_ALIGN.CENTER, space_after=8)
para(tf, False, "intersect / union / difference", size=11, color=GRAY, align=PP_ALIGN.CENTER, space_after=6)
para(tf, False, "两轨解耦媒介", size=11, bold=True, color=ACCENT, align=PP_ALIGN.CENTER, space_after=0)

box = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x3, y, bw, h, ICE, radius=0.06)
tf = box.text_frame
tf.margin_top = Inches(0.12)
para(tf, True, "索引实现多态", size=14.5, bold=True, color=NAVY, space_after=2)
para(tf, False, "BaseIndex（ABC）策略族", size=10.5, color=MIDGRAY, space_after=8)
for t in ["HashIndex：等值倒排 field → value → Set[uid]", "SortedIndex：epoch 有序 + bisect 区间",
          "TimestampNormalizer：多格式 → epoch", "统一 lookup(cond) → CandidateSet"]:
    para(tf, False, "· " + t, size=11.5, color=GRAY, space_after=7)

arrow(s, x1 + bw + 0.05, y + h / 2 - 0.11, gap - 0.1, 0.22, color=MIDGRAY)
arrow(s, x2 + bw + 0.05, y + h / 2 - 0.11, gap - 0.1, 0.22, color=MIDGRAY)

chips_row(s, 5.5, ["值对象隔离", "策略多态", "组合优于继承（不继承“上帝类” SemanticMap）"],
          h=0.62)
box, tf = add_box(s, X0, 6.3, FULL, 0.6)
para(tf, True, "新增对象：QueryConstraints / CandidateSet / BaseConstraintResolver 族 / MetadataIndex + BaseIndex 族 / ConstraintAwarePlanner / ConstraintGroup（职责表见报告 §4.2）",
     size=10.5, bold=False, color=MIDGRAY, space_after=0)


# ================================================================ P10 mechanism 1

s = new_slide()
header(s, "关键机制一：属性与时间索引", 10)
box, tf = add_box(s, X0, 1.45, FULL, 4.4)
bullets = [
    ("路由：", "等值字段 → HashIndex（field → value → Set[uid] 倒排）；时间/数值字段 → SortedIndex（epoch + bisect）"),
    ("选择率：", "SortedIndex 由 bisect 区间计数给出精确选择率；HashIndex 由桶大小直接得出，供 Planner 决策"),
    ("归一化：", "TimestampNormalizer 多格式容错解析（字符串 → epoch float），缺失/非法值走缺省策略，不入索引"),
    ("注册制：", "register_filterable_fields([\"timestamp\", \"speaker\", ...]) 按需启用，避免盲目建索引、控制维护成本"),
    ("维护：", "与 _incremental_aux_retriever_add 并列挂载于写入链；缓存失效沿用 _space_membership_version 版本号模式"),
    ("换页正交：", "L1 换出仅删 payload（保留索引/映射/空间/图）→ 倒排不受 tiered 换出影响，命中后按需 page-in"),
]
for i, (pre, body) in enumerate(bullets):
    para_runs(tf, i == 0, [("● ", False, BLUE), (pre, True, NAVY), (body, False, GRAY)],
              size=13.5, space_after=13)
chips_row(s, 6.0, [("add_unit", None), ("batch_add_units", None), ("delete_unit", None),
                   ("tiered 换入/换出", None)], h=0.55, size=12,
          fill=ICE, gap=0.18)
box, tf = add_box(s, X0, 6.68, FULL, 0.4)
para(tf, True, "▲ 索引增量维护的挂载点（与 BM25/SPLADE 增量并列）", size=10.5, bold=False,
     color=MIDGRAY, align=PP_ALIGN.CENTER, space_after=0)


# ================================================================ P11 mechanism 2

s = new_slide()
header(s, "关键机制二：关系约束谓词", 11)
box, tf = add_box(s, X0, 1.5, FULL, 4.2)
bullets = [
    ("求解方式：", "从实体 E 出发，在 rustworkx 图上做限制深度 BFS（支持方向 / 深度 / 关系类型过滤）→ 可达节点集合"),
    ("谓词同构：", "产出统一为 CandidateSet —— 与属性/时间谓词同构，可直接参与 AND / OR / NOT 组合"),
    ("红线：", "不调用断链 API（search_graph_relations(seed) / get_node_neighbors），直接操作 rx_graph + uid↔index 映射"),
    ("可选增强（暂缓）：", "关系类型倒排 —— 全边扫描在课程规模可接受；无删边 API、维护语义需自行补齐，收益边际递减"),
]
for i, (pre, body) in enumerate(bullets):
    para_runs(tf, i == 0, [("● ", False, BLUE), (pre, True, NAVY), (body, False, GRAY)],
              size=13.5, space_after=14)
# BFS 小示意
chip(s, 1.05, 5.35, 1.1, 0.62, "实体 E", fill=NAVY, color=WHITE, size=13)
arrow(s, 2.25, 5.6, 0.4, 0.14)
chip(s, 2.75, 5.35, 2.6, 0.62, "1-hop 邻居集合", fill=LIGHT, size=12.5)
arrow(s, 5.45, 5.6, 0.4, 0.14)
chip(s, 5.95, 5.35, 2.6, 0.62, "深度限制扩展", fill=LIGHT, size=12.5)
box, tf = add_box(s, 8.6, 5.28, 4.1, 0.8)
para(tf, True, "图上 BFS → uid 集合 →\n与属性谓词统一参与组合", size=11.5, bold=True,
     color=ACCENT, space_after=0)


# ================================================================ P12 mechanism 3 (core)

s = new_slide()
header(s, "关键机制三：候选集代数与自适应调度（核心）", 12)
flow_row(s, 1.5, 1.0, [
    ("① 叶子谓词独立求值", "各 Resolver → CandidateSet"),
    ("② 约束组合", "ConstraintGroup：AND / OR / NOT"),
    ("③ 选择率决策", "|候选| / 总量 < 阈值 ？"),
], sub_size=10.5)
# branch row
chip(s, X0, 2.9, 5.9, 1.05, "pre-filter（是）",
     sub="候选集作为 candidate_uids 下推三路检索器（多态签名现成支持）",
     fill=NAVY, color=WHITE, size=13.5, sub_size=11)
chip(s, X0 + 6.23, 2.9, 5.9, 1.05, "post-filter（否）",
     sub="放大 top_k（复用 max(top_k*3, 50)）→ 融合后置过滤补齐",
     fill=LIGHT, color=NAVY, size=13.5, sub_size=11)
band(s, X0, 4.25, FULL, 0.75,
     "语义相似度只作排序信号最后介入   →   终检：结果硬满足全部约束", size=13)
box, tf = add_box(s, X0, 5.25, FULL, 1.5)
para_runs(tf, True, [("· ", False, BLUE), ("阈值初始 20%，随库规模与后端代价实验标定（不写死）；选择率未知时不冒险下推。", False, GRAY)],
          size=12.5, space_after=7)
para_runs(tf, False, [("· ", False, BLUE), ("空候选 / 候选过小的回退策略可配置：返回空 vs 忽略该约束并告警。", False, GRAY)],
          size=12.5, space_after=7)
para_runs(tf, False, [("· ", False, BLUE), ("为什么是核心：避免全库扫描与重排浪费 —— 语义相似退化为纯排序器，而非过滤器。", False, ACCENT)],
          size=12.5, space_after=0)


# ================================================================ P13 verification story

s = new_slide()
header(s, "大模型辅助核验修正：建议必须过“源码关”", 13)
add_table(s, X0, 1.5, FULL, [
    ["大模型建议", "源码核验事实", "处理"],
    ["需修改三路检索器以支持候选过滤", "三路均已显式支持 candidate_uids（bm25:1499 / splade:959 / semantic_map:2167）",
     "修正（利好）：零修改复用"],
    ["复用现成图 API 实现“实体邻居”", "内部依赖 edge_bfs_search / get_relevant_nodes 等未定义，运行时必抛异常",
     "修正（关键）：自实现 BFS"],
    ["时间字段在 metadata", "实际在 raw_data[“timestamp”]；created_time 默认 None、写入不填充",
     "修正：raw_data 优先 + 规范化器"],
    ["图检索作为第 4 路召回加入融合", "smart_search 入口显式剔除 GRAPH_TRAVERSAL",
     "修正：图仅作约束源"],
    ["自行设计 top_k 扩容策略", "已有 max(top_k*3, 50) 扩容取法（advance_retriever.py:1154-1195）",
     "采纳复用"],
], col_widths=[3.5, 6.0, 2.63], font_size=10.5, header_size=11.5, row_height=0.72)
box, tf = add_box(s, X0, 5.6, FULL, 0.9)
para(tf, True, "方法：建议 → 定位文件/行 → 静态核验（调用链/签名/数据结构）→ 采纳 / 修正 / 弃用，并记录理由。",
     size=12, bold=False, color=GRAY, space_after=4)
para(tf, False, "※ 完整 10 条对照表见备用页 B1 —— 其中 2 条是核验后被动放弃（而非主动取舍）。",
     size=11, color=MIDGRAY, space_after=0)


# ================================================================ P14 comparison

s = new_slide()
header(s, "方案对比与工程判断", 14)
add_table(s, X0, 1.5, 8.35, [
    ["备选思路", "结论", "原因"],
    ["A. smart_search 加 dict 过滤", "未采用", "只能后置过滤，无法利用 pre-filter"],
    ["B. 外部数据库（Milvus/PG）", "未采用", "违背“原生内存、消除跨库 I/O”卖点"],
    ["C. 继承 SemanticMap 派生", "未采用", "组合优于继承，父类已职责过重"],
    ["D. 约束逻辑塞进各 BaseRetriever", "未采用", "重复实现，违反 DRY"],
    ["E. 策略族 Resolver + Planner 调度", "采用", "谓词同构为 CandidateSet，衔接零修改"],
    ["F. 复用既有图 API（断链）", "未采用", "运行时必抛异常；返回边表非节点集"],
], col_widths=[3.35, 1.15, 3.85], font_size=10, header_size=11, row_height=0.62,
    highlight_rows={5})
box, tf = add_box(s, 9.25, 1.5, 3.5, 4.6)
para(tf, True, "五条工程判断", size=14, bold=True, color=NAVY, space_after=8)
for i, t in enumerate([
    "扩展点落在既有签名（candidate_uids）上",
    "先正确性、后性能（P0 基线 → P1 索引）",
    "不信任未核验结论 —— 细节假设错误会让方向跑偏",
    "组合优于继承、门面优于扩散，新职责进新对象",
    "主动裁剪范围：关系类型倒排降级为可选",
]):
    para(tf, False, f"{i+1}. {t}", size=11.5, color=GRAY, space_after=10, line_spacing=1.1)


# ================================================================ P15 plan

s = new_slide()
header(s, "实施计划与预期成果", 15)
# P0/P1 done, P2 doing, P3 next
chip(s, X0, 1.5, 2.95, 0.85, "P0 骨架 + 等价性基线", sub="✓ 已完成", fill=GREEN, color=WHITE, size=12.5, sub_size=10.5)
chip(s, X0 + 3.06, 1.5, 2.95, 0.85, "P1 索引 + pre-filter 下推", sub="✓ 已完成", fill=GREEN, color=WHITE, size=12.5, sub_size=10.5)
chip(s, X0 + 6.12, 1.5, 2.95, 0.85, "P2 关系约束 BFS", sub="进行中", fill=LIGHT, color=NAVY, size=12.5, sub_size=10.5)
chip(s, X0 + 9.18, 1.5, 2.95, 0.85, "P3 调度完善 + 实验", sub="待启动", fill=ICE, color=NAVY, size=12.5, sub_size=10.5)

box, tf = add_box(s, X0, 2.75, FULL, 2.6)
para(tf, True, "实验设计（对应课程要求：比较检索效果、访问开销、维护成本）", size=14, bold=True,
     color=NAVY, space_after=8)
para_runs(tf, False, [("效果：", True, NAVY), ("组合查询测试集上对比 纯语义 / 语义+线性过滤 / 本方案（谓词下推）→ recall@k、nDCG、约束命中率、空结果率", False, GRAY)],
          size=12.5, space_after=7)
para_runs(tf, False, [("开销：", True, NAVY), ("候选集占比 1%–100% 扫描曲线下比较 pre/post-filter → P50/P99 端到端延迟、候选集生成耗时", False, GRAY)],
          size=12.5, space_after=7)
para_runs(tf, False, [("成本：", True, NAVY), ("索引构建时间、增量 add/delete 开销、内存占用（对照 Add Mean 39.7 ms 基线）", False, GRAY)],
          size=12.5, space_after=0)
band(s, X0, 5.6, FULL, 0.85,
     "预期：高选择性约束下 pre-filter 相比“先检索后线性过滤”有数量级延迟优势", size=13.5)
box, tf = add_box(s, X0, 6.65, FULL, 0.45)
para(tf, True, "约束下推 = 用查询优化器的思路重构记忆检索 —— 让语义相似只做它最擅长的事：排序",
     size=12.5, bold=True, color=ACCENT, align=PP_ALIGN.CENTER, space_after=0)


# ================================================================ B1 full table

s = new_slide()
header(s, "备用页 B1：完整核验对照表（10 条）", 16, backup=True)
add_table(s, X0, 1.4, FULL, [
    ["# / 建议", "源码核验事实", "处理"],
    ["1 需修改三路检索器支持候选过滤", "三路均已支持 candidate_uids，**kwargs 透传", "修正（利好）：零修改复用"],
    ["2 filter_memory_units 作约束求解器", "10 操作符但线性扫描；不查 metadata；无分页", "部分采纳：仅作基线与兜底"],
    ["3 时间字段在 metadata", "实际在 raw_data[“timestamp”]；created_time 不填充", "修正：raw_data 优先 + 规范化"],
    ["4 复用 search_graph_relations(seed) / get_node_neighbors", "内部依赖方法未定义，全仓库仅调用点", "修正（关键）：自实现 BFS"],
    ["5 图检索作第 4 路召回", "smart_search 入口剔除 GRAPH_TRAVERSAL", "修正：图仅作约束源"],
    ["6 空间过滤缓存 _get_space_filter_int_ids", "名称不存在；实际为 _space_filter_internal_ids / _get_candidate_uids_set", "修正命名，沿用既有失效模式"],
    ["7 换出破坏索引、需同步维护", "L1 换出仅删 payload，保留索引/映射/空间/图", "采纳事实：倒排与换页正交"],
    ["8 自行设计扩容策略", "已有 max(top_k*3, 50) 扩容取法", "采纳复用"],
    ["9 增量维护需另找挂载点", "写入链末端已有 _incremental_aux_retriever_add 先例", "采纳：并列挂载该模式"],
    ["10 属性倒排作为独立检索源", "需修改 _ensure_retriever_loaded 加载链，非零改动", "不采用：定位为约束层生成器"],
], col_widths=[3.6, 5.9, 2.63], font_size=9, header_size=10, row_height=0.46)


# ================================================================ B2 broken APIs

s = new_slide()
header(s, "备用页 B2：断链 API 细节（关系约束前置风险）", 17, backup=True)
add_table(s, X0, 1.5, FULL, [
    ["调用点（semantic_graph.py）", "被调用方法", "现状"],
    ["search_graph_relations(seed_nodes=...)  :1161", "edge_bfs_search", "未定义（全仓库仅有调用点）"],
    ["search_graph_nodes(search_method=...)  :1126-1132", "hybrid_node_search / _node_similarity_search / _node_fulltext_search", "名称不一致（实际为 _hybrid_node_search 等）"],
    ["get_node_neighbors  :1204", "get_relevant_nodes", "未定义"],
], col_widths=[4.6, 4.9, 2.63], font_size=10, header_size=11, row_height=0.85)
box, tf = add_box(s, X0, 5.0, FULL, 1.5)
para(tf, True, "对策：", size=13.5, bold=True, color=NAVY, space_after=6)
para(tf, False, "· 关系约束不得复用这些断链 API —— Resolver 直接操作 rx_graph + uid↔index 映射；",
     size=12.5, color=GRAY, space_after=5)
para(tf, False, "· 断链修复列为可选任务（P3），作为“分析-发现-修复”素材，不阻塞主线。",
     size=12.5, color=GRAY, space_after=0)


# ================================================================ B3 test plan

s = new_slide()
header(s, "备用页 B3：测试计划矩阵", 18, backup=True)
add_table(s, X0, 1.4, FULL, [
    ["层次", "测试对象", "关键断言"],
    ["单元", "CandidateSet", "集合代数（intersect/union/difference）语义正确、选择性估计合理"],
    ["单元", "MetadataIndex / BaseIndex 族", "增删批一致；未注册不建索引；版本失效重建；Hash/Sorted 同契约"],
    ["单元", "TimestampNormalizer", "多格式 → epoch 容错解析；缺失/非法走缺省策略"],
    ["单元", "各 ConstraintResolver", "属性等值/范围、时间区间、图 BFS（方向/深度/类型）产出正确；空输入边界"],
    ["单元", "ConstraintGroup", "AND/OR/NOT 语义、嵌套组合、空组行为"],
    ["单元", "ConstraintAwarePlanner", "阈值切换 pre/post 决策正确；回退策略生效"],
    ["集成", "search_constrained 端到端", "U1/U2/U3 结果硬满足全部约束；与线性过滤基线等价"],
    ["集成", "向后兼容", "不传约束时行为与现状完全一致"],
    ["回归", "仓库现有测试", "pytest tests/ 全部通过；ruff check 通过"],
], col_widths=[0.95, 3.94, 7.24], font_size=9.5, header_size=10.5, row_height=0.5)
box, tf = add_box(s, X0, 6.5, FULL, 0.5)
para(tf, True, "测试全部离线可跑：Dummy 模型 + monkeypatch，沿用 tests/test_rocksdb_tiered_cache.py:24-81 范式。",
     size=10.5, bold=False, color=MIDGRAY, space_after=0)


# ================================================================ B4 evidence index

s = new_slide()
header(s, "备用页 B4：核验证据索引（行号级）", 19, backup=True)
add_table(s, X0, 1.35, FULL, [
    ["结论", "证据位置"],
    ["candidate_uids 已贯通三路检索器", "bm25_retriever.py:1499-1574；splade_retriever.py:959-997；semantic_map.py:2167"],
    ["smart_search kwargs 透传链与扩容取法", "advance_retriever.py:1154-1195；:937（剔除 GRAPH_TRAVERSAL）"],
    ["filter_memory_units 线性扫描、10 操作符、不查 metadata", "semantic_map.py:3169"],
    ["时间戳在 raw_data[“timestamp”]；created_time 不填充", "locomo_benchmark_episodic.py:462；semantic_map.py:3054"],
    ["SemanticGraph 三处断链 API", "semantic_graph.py:1161 / 1126-1132 / 1204"],
    ["add_relationship 关系维护唯一入口", "semantic_graph.py:610"],
    ["L1 换出仅删 payload、保留索引/映射/空间/图", "semantic_map.py:631-653"],
    ["写入链增量维护挂载点", "semantic_map.py:1217 / 1278 / 1264-1278"],
    ["add_unit / batch_add_units / delete_unit", "semantic_map.py:1280 / 1376 / 1817"],
    ["MemoryUnit 封装的派生视图", "memory_unit.py:10-65"],
    ["SemanticGraph 门面包装模式", "semantic_graph.py:1517-1532"],
    ["benchmark 直连 smart_search", "locomo_benchmark_episodic.py:157"],
    ["测试设施现状与离线范式", "tests/test_rocksdb_tiered_cache.py:24-81"],
    ["三塔编排与量化打包位置", "triple_tower_retriever.py:202、_build_quantification"],
    ["系统架构与分层换页自述", "README_CN.md:52-54、97-99、103-105、121、137-147"],
], col_widths=[5.6, 6.53], font_size=8.5, header_size=9.5, row_height=0.345)


# ================================================================ B5 tiered orthogonality

s = new_slide()
header(s, "备用页 B5：换页与索引正交性（为什么倒排不受换出影响）", 20, backup=True)
flow_row(s, 1.6, 0.9, [
    ("L1（热）", "memory_units 常驻"),
    ("高水位触发", "_trigger_tiered_eviction_if_needed"),
    ("L2 RocksDB（冷）", "仅 payload 换出"),
    ("命中唤回", "get_unit 按需 page-in"),
], sub_size=10.5)
box, tf = add_box(s, X0, 2.9, FULL, 2.9)
bullets = [
    ("换出只做一件事：", "del memory_units[uid]（semantic_map.py:631-653）"),
    ("显式保留：", "FAISS 向量索引、uid ↔ int_id 映射、MemorySpace 成员关系、图拓扑与边"),
    ("对本方案的含义：", "倒排按 uid 存值 → 属性/时间索引与换页正交，不因 payload 换出而失效"),
    ("命中冷 uid：", "get_unit 触发 _add_to_l1_from_tiered_swap 唤回 payload，索引无需重建"),
]
for i, (pre, body) in enumerate(bullets):
    para_runs(tf, i == 0, [("● ", False, BLUE), (pre, True, NAVY), (body, False, GRAY)],
              size=13.5, space_after=13)
band(s, X0, 6.0, FULL, 0.75, "结论：索引生命周期独立于 payload 分层 —— 无需在索引侧做换页同步维护",
     size=13)


# ---------------------------------------------------------------- save

prs.save(OUTPUT)
print(f"Saved: {OUTPUT} | slides: {len(prs.slides)}")